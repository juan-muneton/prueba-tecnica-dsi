"""End: arma el Excel diario desde la cola y escribe el resumen de la ejecución.

Equivale al estado End de mi plantilla (consolidar el Excel y el resumen).
- El Excel se reconstruye completo en cada ejecución (C3): una fila por id (D1),
  en el orden del archivo, con las 8 columnas de D3. Por eso nunca duplica.
- Se escribe en un temporal y luego se reemplaza; si el Excel está abierto y
  bloqueado, se detiene con EXCEL_BLOQUEADO y la cola queda intacta.
- El resumen (D4) va al log. El código de salida es 0 si no quedan errores de
  sistema y 2 si quedan pendientes (para que el Programador de tareas avise).
"""
from __future__ import annotations

import logging
import os
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from framework.exceptions import SystemException
from framework.transaction_queue import TransactionQueue
from interactions.dispatcher import DispatcherResult

LOGGER = logging.getLogger("consulta_cuentas")

COLUMNS = ("ID Cuenta", "Líneas", "User ID", "Título", "Contenido", "Fecha de consulta", "Estado", "Observación")
WIDTHS = (12, 14, 9, 45, 60, 20, 18, 60)
EXIT_OK, EXIT_PENDING = 0, 2


def build_rows(config: dict[str, Any], batch: DispatcherResult, queue: TransactionQueue) -> list[tuple]:
    """Una fila por transacción, en el orden del archivo. Las claves de estado y
    observación se traducen aquí con el catálogo de config.json."""
    rows = []
    for t in batch.transactions:
        lines = ", ".join(map(str, t.lines))
        if not t.is_valid:
            state, code, record = "EstadoNegocio", t.invalid_code, {}
            identifier: Any = t.raw_value
        else:
            record = queue.record(t.account_id) or {}
            state = record.get("estado", "EstadoSistema")
            code = record.get("codigo", "PENDIENTE_POR_CORTE")
            identifier = t.account_id
        rows.append((
            identifier,
            lines,
            record.get("user_id"),
            record.get("title"),
            record.get("body"),
            record.get("fecha_consulta"),
            config[state],
            config["Observaciones"][code],
        ))
    return rows


def write_excel(path: Path, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    tmp = path.with_name(f"~tmp_{path.name}")
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise SystemException("EXCEL_BLOQUEADO", f"No se pudo reemplazar {path.name}: está abierto en otro programa") from exc


def log_summary(config: dict[str, Any], batch: DispatcherResult, rows: list[tuple], consulted_now: int, resolved_before: int) -> int:
    state_col, obs_col = COLUMNS.index("Estado"), COLUMNS.index("Observación")
    failed = Counter((row[state_col], row[obs_col]) for row in rows if row[state_col] != config["EstadoExitoso"])
    resolved = len(rows) - sum(failed.values())

    LOGGER.info("===== RESUMEN DE LA EJECUCIÓN =====")
    LOGGER.info("Filas en el Excel: %s (de %s líneas del archivo)", len(rows), batch.total_lines)
    LOGGER.info("Consultadas en esta ejecución: %s · ya resueltas antes: %s", consulted_now, resolved_before)
    LOGGER.info("Resueltas: %s", resolved)
    LOGGER.info("Fallidas: %s", sum(failed.values()))
    for (state, observation), count in sorted(failed.items(), key=lambda item: (item[0][0], -item[1], item[0][1])):
        LOGGER.info("  %s · %s: %s", state, observation, count)

    pending = sum(count for (state, _), count in failed.items() if state == config["EstadoSistema"])
    return EXIT_PENDING if pending else EXIT_OK


def end_process(config: dict[str, Any], batch: DispatcherResult, queue: TransactionQueue, consulted_now: int) -> int:
    rows = build_rows(config, batch, queue)
    path = config["CarpetaOutput"] / queue.output_name
    write_excel(path, rows)
    LOGGER.info("Excel generado; archivo=%s filas=%s", path.name, len(rows))
    return log_summary(config, batch, rows, consulted_now, queue.resolved_before)
