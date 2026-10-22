"""Arnés del paso 5: el proceso completo (Main, Process y End) con una API falsa. Sin red.

Cada prueba arma un proyecto temporal (Configuracion + Funcional/Input) con el
config.json y el CSV reales, y corre Main de principio a fin.

Ejecutar desde Fuentes/ConsultaCuentas:  python -m unittest discover -s tests -t .
"""
import io
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import main
from framework.init_all_settings import DEFAULT_CONFIG_PATH, LOGGER, load_config
from libs.jsonplaceholder_lib import HttpResponse

REAL_CSV = load_config()["ArchivoEntrada"]
EXCEL = "Funcional/Output/resultado_cuentas_{}.xlsx"


class FakeApi:
    """Se comporta como jsonplaceholder: posts del 1 al 100; el resto da 404."""

    def __init__(self, down=False, crash_after=None):
        self.down = down
        self.crash_after = crash_after
        self.calls = []

    def get(self, url, timeout):
        self.calls.append(url)
        if self.crash_after is not None and len(self.calls) > self.crash_after:
            raise KeyboardInterrupt  # simula que se apaga la máquina a mitad
        if self.down:
            raise socket.timeout()
        account_id = int(url.rsplit("/", 1)[1])
        if account_id > 100:
            return HttpResponse(404, b"{}")
        body = {"userId": account_id % 10 + 1, "id": account_id, "title": f"titulo {account_id}", "body": "cuerpo"}
        return HttpResponse(200, json.dumps(body).encode("utf-8"))


class MainTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / "Configuracion").mkdir()
        (self.root / "Funcional" / "Input").mkdir(parents=True)
        shutil.copy(DEFAULT_CONFIG_PATH, self.root / "Configuracion" / "config.json")
        self.csv = self.root / "Funcional" / "Input" / "cuentas.csv"
        shutil.copy(REAL_CSV, self.csv)

    def tearDown(self):
        for handler in list(LOGGER.handlers):
            LOGGER.removeHandler(handler)
            handler.close()
        self.tmp.cleanup()

    def run_main(self, api):
        screen = io.StringIO()
        with redirect_stdout(screen):
            code = main.run(self.root / "Configuracion" / "config.json", api, sleep=lambda s: None)
        return code, screen.getvalue()

    def excel_rows(self):
        [path] = (self.root / "Funcional" / "Output").glob("resultado_cuentas_*.xlsx")
        ws = load_workbook(path).active
        return [tuple(c.value for c in row) for row in ws.iter_rows()]

    def test_full_run_with_the_challenge_csv(self):
        api = FakeApi()
        code, screen = self.run_main(api)
        self.assertEqual(code, 0)
        self.assertEqual(len(api.calls), 15)  # 15 ids únicos, los repetidos una sola vez
        rows = self.excel_rows()
        self.assertEqual(rows[0], ("ID Cuenta", "Líneas", "User ID", "Título", "Contenido", "Fecha de consulta", "Estado", "Observación"))
        self.assertEqual(len(rows), 19)  # encabezado + 18 filas
        by_id = {row[0]: row for row in rows[1:]}
        self.assertEqual(by_id[7][1], "2, 6, 18")
        self.assertEqual(by_id[7][2:5], (8, "titulo 7", "cuerpo"))
        self.assertEqual(by_id[7][6:], ("Exitoso", "Cuenta encontrada."))
        self.assertEqual(by_id[101][6:], ("Regla de Negocio", "La cuenta no existe en el sistema."))
        self.assertEqual(by_id["abc"][1:], ("9", None, None, None, None, "Regla de Negocio", "El valor no es un número de cuenta."))
        self.assertEqual(by_id[None][6:], ("Regla de Negocio", "La línea está vacía."))
        self.assertEqual(by_id["0"][6:], ("Regla de Negocio", "El número de cuenta debe ser mayor que cero."))
        self.assertIn("Consultadas en esta ejecución: 15 · ya resueltas antes: 0", screen)
        self.assertIn("Resueltas: 13", screen)
        self.assertIn("Fallidas: 5", screen)
        self.assertIn("Regla de Negocio · La cuenta no existe en el sistema.: 2", screen)

    def test_second_run_queries_nothing_and_does_not_duplicate(self):
        self.run_main(FakeApi())
        first = self.excel_rows()
        api = FakeApi()
        code, screen = self.run_main(api)
        self.assertEqual((code, api.calls), (0, []))
        self.assertEqual(self.excel_rows(), first)
        self.assertIn("Consultadas en esta ejecución: 0 · ya resueltas antes: 15", screen)

    def test_crash_in_the_middle_resumes_where_it_stopped(self):
        with self.assertRaises(KeyboardInterrupt):
            self.run_main(FakeApi(crash_after=4))
        api = FakeApi()
        code, screen = self.run_main(api)
        self.assertEqual(code, 0)
        self.assertEqual(len(api.calls), 11)  # 4 ya guardados antes de la caída
        self.assertIn("ya resueltas antes: 4", screen)
        self.assertEqual(len(self.excel_rows()), 19)

    def test_api_down_cuts_after_three_and_next_run_recovers(self):
        api = FakeApi(down=True)
        code, screen = self.run_main(api)
        self.assertEqual(code, 2)
        self.assertEqual(len(api.calls), 12)  # 3 ids x 4 intentos, después se corta
        self.assertIn("API caída, se detienen las consultas; errores_seguidos=3 pendientes=12", screen)
        states = [row[7] for row in self.excel_rows()[1:]]
        self.assertEqual(states.count("El sistema no respondió a tiempo; se agotaron los reintentos."), 3)
        self.assertEqual(states.count("No se consultó porque el sistema estaba caído; se intentará en la siguiente ejecución."), 12)

        api = FakeApi()
        code, screen = self.run_main(api)
        self.assertEqual((code, len(api.calls)), (0, 15))
        self.assertIn("Resueltas: 13", screen)

    def test_locked_excel_stops_and_keeps_the_queue(self):
        real_replace = os.replace

        def excel_locked(src, dst):
            # Solo el Excel está bloqueado; la cola se guarda normal
            if str(dst).endswith(".xlsx"):
                raise PermissionError(13, "El archivo está abierto en otro programa")
            return real_replace(src, dst)

        with mock.patch("os.replace", side_effect=excel_locked):
            code, screen = self.run_main(FakeApi())
        self.assertEqual(code, 1)
        self.assertIn("codigo=EXCEL_BLOQUEADO", screen)
        self.assertEqual(list((self.root / "Funcional" / "Output").iterdir()), [])  # sin temporales
        api = FakeApi()
        code, _ = self.run_main(api)
        self.assertEqual((code, api.calls), (0, []))  # la cola quedó guardada

    def test_only_header_ends_without_input(self):
        self.csv.write_text("id_cuenta\n", encoding="utf-8")
        api = FakeApi()
        code, screen = self.run_main(api)
        self.assertEqual((code, api.calls), (0, []))
        self.assertIn("Proceso terminado sin insumo", screen)
        self.assertEqual(list((self.root / "Funcional" / "Output").glob("*.xlsx")), [])

    def test_missing_input_file_stops(self):
        self.csv.unlink()
        code, screen = self.run_main(FakeApi())
        self.assertEqual(code, 1)
        self.assertIn("codigo=ENTRADA_NO_EXISTE", screen)

    def test_config_error_on_redirected_screen_keeps_accents(self):
        # Como lo lanza el Programador de tareas: la salida de errores va a un archivo
        config = self.root / "Configuracion" / "config.json"
        config.write_text('{"Settings": {}}', encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(Path(main.__file__)), "--config", str(config)],
            capture_output=True,
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("Falta la sección 'Constants'", result.stderr.decode("utf-8"))

    def test_invalid_config_stops_before_anything(self):
        (self.root / "Configuracion" / "config.json").write_text("{}", encoding="utf-8")
        with mock.patch("sys.stderr", new_callable=io.StringIO) as err:
            code, _ = self.run_main(FakeApi())
        self.assertEqual(code, 1)
        self.assertIn("codigo=CONFIG_INVALIDA", err.getvalue())


if __name__ == "__main__":
    unittest.main()
